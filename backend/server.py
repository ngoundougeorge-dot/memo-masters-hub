from dotenv import load_dotenv
load_dotenv()

import os
import uuid
import secrets
import asyncio
import logging
from io import BytesIO
from datetime import datetime, timezone, timedelta

import bcrypt
import jwt
import requests
import resend
from fastapi import FastAPI, APIRouter, HTTPException, Request, Depends, UploadFile, File, Form, Query, Header
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse, Response
from pydantic import BaseModel, EmailStr, Field
from motor.motor_asyncio import AsyncIOMotorClient
from typing import Optional, List

# ---------------------------------------------------------------------------
# Config
# ---------------------------------------------------------------------------
MONGO_URL = os.environ["MONGO_URL"]
DB_NAME = os.environ["DB_NAME"]
JWT_SECRET = os.environ["JWT_SECRET"]
JWT_ALGORITHM = "HS256"
WRITER_ACCESS_KEY = os.environ.get("WRITER_ACCESS_KEY", "")
EMERGENT_LLM_KEY = os.environ.get("EMERGENT_LLM_KEY", "")
APP_NAME = "memoirepro"

STORAGE_BASE = (os.environ.get("INTEGRATION_PROXY_URL") or "").strip() or "https://integrations.emergentagent.com"
STORAGE_URL = STORAGE_BASE.rstrip("/") + "/objstore/api/v1/storage"

RESEND_API_KEY = os.environ.get("RESEND_API_KEY", "").strip()
SENDER_EMAIL = os.environ.get("SENDER_EMAIL", "onboarding@resend.dev")
WRITER_NOTIFY_EMAIL = os.environ.get("WRITER_NOTIFY_EMAIL", "").strip()

logger = logging.getLogger("memoirepro")
logger.setLevel(logging.INFO)
if not logger.handlers:
    _h = logging.StreamHandler()
    _h.setFormatter(logging.Formatter("%(levelname)s:memoirepro:%(message)s"))
    logger.addHandler(_h)
    logger.propagate = False

client = AsyncIOMotorClient(MONGO_URL)
db = client[DB_NAME]

app = FastAPI(title="MémoirePro API")
api = APIRouter(prefix="/api")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=False,
    allow_methods=["*"],
    allow_headers=["*"],
)

# ---------------------------------------------------------------------------
# Object storage
# ---------------------------------------------------------------------------
storage_key = None

MIME_TYPES = {
    "jpg": "image/jpeg", "jpeg": "image/jpeg", "png": "image/png",
    "gif": "image/gif", "webp": "image/webp", "pdf": "application/pdf",
    "doc": "application/msword", "txt": "text/plain",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "odt": "application/vnd.oasis.opendocument.text",
}


def init_storage(force: bool = False):
    global storage_key
    if storage_key and not force:
        return storage_key
    resp = requests.post(f"{STORAGE_URL}/init", json={"emergent_key": EMERGENT_LLM_KEY}, timeout=30)
    resp.raise_for_status()
    storage_key = resp.json()["storage_key"]
    return storage_key


def put_object(path: str, data: bytes, content_type: str) -> dict:
    key = init_storage()
    resp = requests.put(
        f"{STORAGE_URL}/objects/{path}",
        headers={"X-Storage-Key": key, "Content-Type": content_type},
        data=data, timeout=120,
    )
    if resp.status_code == 404:
        key = init_storage(force=True)
        resp = requests.put(
            f"{STORAGE_URL}/objects/{path}",
            headers={"X-Storage-Key": key, "Content-Type": content_type},
            data=data, timeout=120,
        )
    resp.raise_for_status()
    return resp.json()


def get_object(path: str):
    key = init_storage()
    resp = requests.get(f"{STORAGE_URL}/objects/{path}", headers={"X-Storage-Key": key}, timeout=60)
    if resp.status_code == 404:
        key = init_storage(force=True)
        resp = requests.get(f"{STORAGE_URL}/objects/{path}", headers={"X-Storage-Key": key}, timeout=60)
    resp.raise_for_status()
    return resp.content, resp.headers.get("Content-Type", "application/octet-stream")


# ---------------------------------------------------------------------------
# Auth helpers
# ---------------------------------------------------------------------------
def hash_password(password: str) -> str:
    return bcrypt.hashpw(password.encode("utf-8"), bcrypt.gensalt()).decode("utf-8")


def verify_password(plain: str, hashed: str) -> bool:
    try:
        return bcrypt.checkpw(plain.encode("utf-8"), hashed.encode("utf-8"))
    except Exception:
        return False


def create_access_token(user_id: str, email: str) -> str:
    payload = {"sub": user_id, "email": email, "exp": datetime.now(timezone.utc) + timedelta(days=7), "type": "access"}
    return jwt.encode(payload, JWT_SECRET, algorithm=JWT_ALGORITHM)


async def get_current_user(request: Request) -> dict:
    auth_header = request.headers.get("Authorization", "")
    token = auth_header[7:] if auth_header.startswith("Bearer ") else None
    if not token:
        raise HTTPException(status_code=401, detail="Non authentifié")
    try:
        payload = jwt.decode(token, JWT_SECRET, algorithms=[JWT_ALGORITHM])
    except jwt.ExpiredSignatureError:
        raise HTTPException(status_code=401, detail="Session expirée")
    except jwt.InvalidTokenError:
        raise HTTPException(status_code=401, detail="Jeton invalide")
    user = await db.users.find_one({"id": payload["sub"]})
    if not user:
        raise HTTPException(status_code=401, detail="Utilisateur introuvable")
    user.pop("password_hash", None)
    user.pop("_id", None)
    return user


async def optional_user(request: Request):
    try:
        return await get_current_user(request)
    except HTTPException:
        return None


def now_iso():
    return datetime.now(timezone.utc).isoformat()


# ---------------------------------------------------------------------------
# Email notifications (Resend, graceful fallback to log)
# ---------------------------------------------------------------------------
STATUS_FR = {
    "nouveau": "Paiement en attente", "paiement_recu": "Paiement confirmé",
    "documents_envoyes": "Documents reçus", "en_cours": "En préparation",
    "redaction": "Rédaction en cours", "livre": "Document livré",
}


def _email_html(title: str, body: str, footer: str = "") -> str:
    return f"""
    <div style="font-family:Inter,Arial,sans-serif;background:#f7f4ec;padding:24px">
      <table width="100%" cellpadding="0" cellspacing="0" style="max-width:560px;margin:0 auto;background:#ffffff;border-radius:12px;overflow:hidden;border:1px solid #e6e0d3">
        <tr><td style="background:#1b3a2d;padding:20px 24px">
          <span style="color:#c99a3f;font-size:20px;font-weight:700;font-family:Georgia,serif">MémoirePro</span>
        </td></tr>
        <tr><td style="padding:28px 24px">
          <h1 style="margin:0 0 12px;font-size:20px;color:#1b3a2d;font-family:Georgia,serif">{title}</h1>
          <div style="font-size:15px;line-height:1.7;color:#333">{body}</div>
          {f'<p style="margin-top:20px;font-size:13px;color:#888">{footer}</p>' if footer else ''}
        </td></tr>
        <tr><td style="background:#f7f4ec;padding:16px 24px;font-size:12px;color:#999;text-align:center">
          © 2026 MémoirePro — Rédaction académique pour l'Afrique francophone.
        </td></tr>
      </table>
    </div>"""


async def send_email(to: str, subject: str, html: str):
    if not to:
        return
    if not RESEND_API_KEY:
        logger.info(f"[email:simulation] to={to} subject={subject!r} (RESEND_API_KEY absent)")
        return
    try:
        resend.api_key = RESEND_API_KEY
        params = {"from": SENDER_EMAIL, "to": [to], "subject": subject, "html": html}
        await asyncio.to_thread(resend.Emails.send, params)
        logger.info(f"[email:sent] to={to} subject={subject!r}")
    except Exception as e:
        logger.error(f"[email:error] to={to} err={e}")


async def notify_status_change(order: dict, new_status: str):
    label = STATUS_FR.get(new_status, new_status)
    ref = order["id"][:8]
    subject_client = f"MémoirePro — Commande {ref} : {label}"
    body_client = (
        f"Bonjour {order.get('full_name', '')},<br><br>"
        f"Le statut de votre commande <strong>{ref}</strong> « {order.get('subject','')} » "
        f"vient de passer à : <strong>{label}</strong>.<br><br>"
        f"Vous pouvez suivre l'avancement dans votre espace client."
    )
    await send_email(order.get("email"), subject_client,
                     _email_html("Mise à jour de votre commande", body_client,
                                 "Conservez votre numéro de commande pour tout suivi."))
    if WRITER_NOTIFY_EMAIL:
        body_writer = (
            f"La commande <strong>{ref}</strong> ({order.get('full_name','')} · {order.get('email','')}) "
            f"est désormais au statut : <strong>{label}</strong>.<br>"
            f"Sujet : {order.get('subject','')}"
        )
        await send_email(WRITER_NOTIFY_EMAIL, f"[Rédaction] {ref} → {label}",
                         _email_html("Mise à jour commande (rédacteur)", body_writer))


async def notify_documents_submitted(order: dict):
    if WRITER_NOTIFY_EMAIL:
        ref = order["id"][:8]
        body = (
            f"{order.get('full_name','')} ({order.get('email','')}) vient de soumettre "
            f"{len(order.get('file_ids', []))} document(s) pour la commande <strong>{ref}</strong>.<br>"
            f"Sujet : {order.get('subject','')}"
        )
        await send_email(WRITER_NOTIFY_EMAIL, f"[Rédaction] Nouveaux documents — {ref}",
                         _email_html("Documents soumis par un client", body))


# ---------------------------------------------------------------------------
# Models
# ---------------------------------------------------------------------------
class RegisterInput(BaseModel):
    name: str = Field(min_length=2, max_length=120)
    email: EmailStr
    password: str = Field(min_length=6, max_length=200)


class LoginInput(BaseModel):
    email: EmailStr
    password: str


class OrderInput(BaseModel):
    full_name: str = Field(min_length=2, max_length=120)
    email: EmailStr
    phone: str = Field(min_length=6, max_length=40)
    document_type: str
    academic_level: Optional[str] = None
    subject: str = Field(min_length=5, max_length=2000)
    instructions: Optional[str] = None
    deadline: Optional[str] = None
    pages: Optional[int] = None
    payment_method: Optional[str] = None
    file_ids: List[str] = []


class AttachFilesInput(BaseModel):
    order_id: str
    file_ids: List[str]


class SubmitDocsInput(BaseModel):
    order_id: str


class WriterKeyInput(BaseModel):
    key: Optional[str] = None


class MarkSeenInput(BaseModel):
    key: Optional[str] = None
    order_ids: Optional[List[str]] = None
    all: Optional[bool] = False


class StatusUpdateInput(BaseModel):
    key: Optional[str] = None
    order_id: str
    status: str
    price_fcfa: Optional[int] = None


class DocumentInput(BaseModel):
    title: str = "Sans titre"
    content: str = ""


class AIChatInput(BaseModel):
    message: str
    document_id: Optional[str] = None
    context: Optional[str] = None
    mode: Optional[str] = "chat"  # chat | improve | outline | expand


def clean(doc: dict) -> dict:
    if doc:
        doc.pop("_id", None)
    return doc


# ---------------------------------------------------------------------------
# Auth routes
# ---------------------------------------------------------------------------
@api.post("/auth/register")
async def register(data: RegisterInput):
    email = data.email.lower().strip()
    if await db.users.find_one({"email": email}):
        raise HTTPException(status_code=400, detail="Un compte existe déjà avec cet email")
    uid = str(uuid.uuid4())
    doc = {
        "id": uid, "email": email, "password_hash": hash_password(data.password),
        "name": data.name.strip(), "role": "client", "created_at": now_iso(),
    }
    await db.users.insert_one(doc)
    token = create_access_token(uid, email)
    return {"token": token, "user": {"id": uid, "email": email, "name": doc["name"], "role": "client"}}


@api.post("/auth/login")
async def login(data: LoginInput):
    email = data.email.lower().strip()
    user = await db.users.find_one({"email": email})
    if not user or not verify_password(data.password, user["password_hash"]):
        raise HTTPException(status_code=401, detail="Email ou mot de passe incorrect")
    token = create_access_token(user["id"], email)
    return {"token": token, "user": {"id": user["id"], "email": email, "name": user["name"], "role": user.get("role", "client")}}


@api.get("/auth/me")
async def me(user: dict = Depends(get_current_user)):
    return {"user": user}


# ---------------------------------------------------------------------------
# File upload
# ---------------------------------------------------------------------------
@api.post("/files/upload")
async def upload_file(file: UploadFile = File(...), user=Depends(optional_user)):
    data = await file.read()
    if len(data) > 15 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="Fichier trop volumineux (max 15 Mo)")
    ext = file.filename.rsplit(".", 1)[-1].lower() if "." in file.filename else "bin"
    fid = str(uuid.uuid4())
    path = f"{APP_NAME}/uploads/{fid}.{ext}"
    content_type = file.content_type or MIME_TYPES.get(ext, "application/octet-stream")
    try:
        result = put_object(path, data, content_type)
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Échec de l'envoi : {e}")
    rec = {
        "id": fid, "storage_path": result["path"], "original_filename": file.filename,
        "content_type": content_type, "size": result.get("size", len(data)),
        "is_deleted": False, "created_at": now_iso(),
    }
    await db.files.insert_one(rec)
    return {"id": fid, "filename": file.filename, "size": rec["size"]}


@api.get("/files/{file_id}/download")
async def download_file(file_id: str):
    rec = await db.files.find_one({"id": file_id, "is_deleted": False})
    if not rec:
        raise HTTPException(status_code=404, detail="Fichier introuvable")
    data, ct = get_object(rec["storage_path"])
    return Response(
        content=data, media_type=rec.get("content_type", ct),
        headers={"Content-Disposition": f'inline; filename="{rec["original_filename"]}"'},
    )


# ---------------------------------------------------------------------------
# Orders
# ---------------------------------------------------------------------------
@api.post("/orders")
async def create_order(data: OrderInput, user=Depends(optional_user)):
    oid = str(uuid.uuid4())
    doc = {
        "id": oid, "user_id": user["id"] if user else None,
        "full_name": data.full_name, "email": data.email.lower(), "phone": data.phone,
        "document_type": data.document_type, "academic_level": data.academic_level,
        "subject": data.subject, "instructions": data.instructions, "deadline": data.deadline,
        "pages": data.pages, "price_fcfa": None, "payment_method": data.payment_method,
        "file_ids": data.file_ids or [], "status": "nouveau",
        "documents_submitted_at": None, "writer_seen_at": None, "created_at": now_iso(),
    }
    await db.orders.insert_one(doc)
    return {"id": oid}


def order_public(o: dict) -> dict:
    return {
        "id": o["id"], "status": o["status"], "subject": o["subject"],
        "document_type": o["document_type"], "created_at": o["created_at"],
        "price_fcfa": o.get("price_fcfa"), "pages": o.get("pages"), "deadline": o.get("deadline"),
        "payment_method": o.get("payment_method"), "instructions": o.get("instructions"),
        "file_ids": o.get("file_ids", []), "documents_submitted_at": o.get("documents_submitted_at"),
    }


@api.get("/orders/{order_id}")
async def get_order(order_id: str):
    o = await db.orders.find_one({"id": order_id})
    if not o:
        raise HTTPException(status_code=404, detail="Commande introuvable")
    return order_public(o)


@api.get("/orders")
async def my_orders(user: dict = Depends(get_current_user)):
    cursor = db.orders.find({"$or": [{"user_id": user["id"]}, {"email": user["email"]}]}).sort("created_at", -1)
    return [order_public(o) async for o in cursor]


@api.post("/orders/attach-files")
async def attach_files(data: AttachFilesInput):
    o = await db.orders.find_one({"id": data.order_id})
    if not o:
        raise HTTPException(status_code=404, detail="Commande introuvable")
    if o.get("documents_submitted_at"):
        raise HTTPException(status_code=400, detail="Les documents ont déjà été soumis et sont verrouillés.")
    merged = list(dict.fromkeys([*o.get("file_ids", []), *data.file_ids]))[:40]
    await db.orders.update_one({"id": data.order_id}, {"$set": {"file_ids": merged}})
    return {"file_ids": merged}


@api.post("/orders/submit-documents")
async def submit_documents(data: SubmitDocsInput):
    o = await db.orders.find_one({"id": data.order_id})
    if not o:
        raise HTTPException(status_code=404, detail="Commande introuvable")
    if o.get("documents_submitted_at"):
        return {"ok": True, "already": True}
    if not o.get("file_ids"):
        raise HTTPException(status_code=400, detail="Ajoutez au moins un document avant de soumettre.")
    await db.orders.update_one(
        {"id": data.order_id},
        {"$set": {"status": "documents_envoyes", "documents_submitted_at": now_iso()}},
    )
    o["status"] = "documents_envoyes"
    asyncio.create_task(notify_documents_submitted(o))
    asyncio.create_task(notify_status_change(o, "documents_envoyes"))
    return {"ok": True, "already": False}


# ---------------------------------------------------------------------------
# Writer / redacteur
# ---------------------------------------------------------------------------
def check_writer(key: Optional[str], user: Optional[dict]):
    if user and user.get("role") == "admin":
        return
    if key and WRITER_ACCESS_KEY and key == WRITER_ACCESS_KEY:
        return
    raise HTTPException(status_code=403, detail="Accès refusé")


@api.post("/writer/orders")
async def writer_orders(data: WriterKeyInput, user=Depends(optional_user)):
    check_writer(data.key, user)
    cursor = db.orders.find({}).sort([("documents_submitted_at", -1), ("created_at", -1)]).limit(100)
    rows = []
    async for o in cursor:
        rows.append({
            "id": o["id"], "full_name": o["full_name"], "email": o["email"], "subject": o["subject"],
            "document_type": o["document_type"], "status": o["status"], "price_fcfa": o.get("price_fcfa"),
            "pages": o.get("pages"), "deadline": o.get("deadline"), "created_at": o["created_at"],
            "documents_submitted_at": o.get("documents_submitted_at"), "file_ids": o.get("file_ids", []),
            "writer_seen_at": o.get("writer_seen_at"),
        })
    return rows


@api.post("/writer/mark-seen")
async def writer_mark_seen(data: MarkSeenInput, user=Depends(optional_user)):
    check_writer(data.key, user)
    now = now_iso()
    if data.all:
        res = await db.orders.update_many(
            {"status": "documents_envoyes", "writer_seen_at": None}, {"$set": {"writer_seen_at": now}}
        )
        return {"ok": True, "count": res.modified_count}
    if data.order_ids:
        res = await db.orders.update_many({"id": {"$in": data.order_ids}}, {"$set": {"writer_seen_at": now}})
        return {"ok": True, "count": res.modified_count}
    raise HTTPException(status_code=400, detail="Aucune commande à marquer.")


@api.post("/writer/update-status")
async def writer_update_status(data: StatusUpdateInput, user=Depends(optional_user)):
    check_writer(data.key, user)
    o = await db.orders.find_one({"id": data.order_id})
    if not o:
        raise HTTPException(status_code=404, detail="Commande introuvable")
    update = {"status": data.status}
    if data.price_fcfa is not None:
        update["price_fcfa"] = data.price_fcfa
    await db.orders.update_one({"id": data.order_id}, {"$set": update})
    if o.get("status") != data.status:
        o.update(update)
        asyncio.create_task(notify_status_change(o, data.status))
    return {"ok": True}


# ---------------------------------------------------------------------------
# Editor documents
# ---------------------------------------------------------------------------
@api.get("/documents")
async def list_documents(user: dict = Depends(get_current_user)):
    cursor = db.documents.find({"user_id": user["id"]}).sort("updated_at", -1)
    return [clean(d) async for d in cursor]


@api.post("/documents")
async def create_document(data: DocumentInput, user: dict = Depends(get_current_user)):
    doc = {
        "id": str(uuid.uuid4()), "user_id": user["id"], "title": data.title or "Sans titre",
        "content": data.content or "", "created_at": now_iso(), "updated_at": now_iso(),
    }
    await db.documents.insert_one(doc)
    return clean(doc)


@api.get("/documents/{doc_id}")
async def get_document(doc_id: str, user: dict = Depends(get_current_user)):
    d = await db.documents.find_one({"id": doc_id, "user_id": user["id"]})
    if not d:
        raise HTTPException(status_code=404, detail="Document introuvable")
    return clean(d)


@api.put("/documents/{doc_id}")
async def update_document(doc_id: str, data: DocumentInput, user: dict = Depends(get_current_user)):
    res = await db.documents.update_one(
        {"id": doc_id, "user_id": user["id"]},
        {"$set": {"title": data.title, "content": data.content, "updated_at": now_iso()}},
    )
    if res.matched_count == 0:
        raise HTTPException(status_code=404, detail="Document introuvable")
    d = await db.documents.find_one({"id": doc_id})
    return clean(d)


@api.delete("/documents/{doc_id}")
async def delete_document(doc_id: str, user: dict = Depends(get_current_user)):
    await db.documents.delete_one({"id": doc_id, "user_id": user["id"]})
    return {"ok": True}


def _safe_filename(title: str, ext: str) -> str:
    base = "".join(c if c.isalnum() or c in " -_" else "_" for c in (title or "memoire")).strip() or "memoire"
    return f"{base[:60]}.{ext}"


def _build_docx(title: str, content: str) -> bytes:
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(12)

    h = doc.add_heading(title or "Mémoire", level=0)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for block in (content or "").split("\n"):
        line = block.rstrip()
        if not line:
            doc.add_paragraph("")
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:], level=3)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        elif line.startswith("# "):
            doc.add_heading(line[2:], level=1)
        elif line.lstrip().startswith(("- ", "* ")):
            doc.add_paragraph(line.lstrip()[2:], style="List Bullet")
        else:
            p = doc.add_paragraph(line)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _build_pdf(title: str, content: str) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.lib.enums import TA_JUSTIFY, TA_CENTER
    from reportlab.lib.colors import HexColor
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=2.5 * cm, bottomMargin=2.5 * cm,
                            leftMargin=2.5 * cm, rightMargin=2.5 * cm)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("Title2", parent=styles["Title"], textColor=HexColor("#1b3a2d"), alignment=TA_CENTER, fontSize=22, spaceAfter=24)
    body = ParagraphStyle("Body2", parent=styles["Normal"], fontSize=12, leading=18, alignment=TA_JUSTIFY, spaceAfter=8)
    h1 = ParagraphStyle("H1", parent=styles["Heading1"], textColor=HexColor("#1b3a2d"), fontSize=16)
    h2 = ParagraphStyle("H2", parent=styles["Heading2"], textColor=HexColor("#1b3a2d"), fontSize=14)

    def esc(s):
        return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    flow = [Paragraph(esc(title or "Mémoire"), title_style)]
    for line in (content or "").split("\n"):
        line = line.rstrip()
        if not line:
            flow.append(Spacer(1, 8))
        elif line.startswith("## "):
            flow.append(Paragraph(esc(line[3:]), h2))
        elif line.startswith("# "):
            flow.append(Paragraph(esc(line[2:]), h1))
        else:
            flow.append(Paragraph(esc(line), body))
    doc.build(flow)
    return buf.getvalue()


@api.get("/documents/{doc_id}/export")
async def export_document(doc_id: str, format: str = Query("pdf"), user: dict = Depends(get_current_user)):
    d = await db.documents.find_one({"id": doc_id, "user_id": user["id"]})
    if not d:
        raise HTTPException(status_code=404, detail="Document introuvable")
    title, content = d.get("title", "Mémoire"), d.get("content", "")
    if format == "docx":
        data = await asyncio.to_thread(_build_docx, title, content)
        media = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        fname = _safe_filename(title, "docx")
    else:
        data = await asyncio.to_thread(_build_pdf, title, content)
        media = "application/pdf"
        fname = _safe_filename(title, "pdf")
    return Response(content=data, media_type=media,
                    headers={"Content-Disposition": f'attachment; filename="{fname}"'})


# ---------------------------------------------------------------------------
# AI assistant (Claude Sonnet 4.6, streaming SSE)
# ---------------------------------------------------------------------------
SYSTEM_PROMPTS = {
    "chat": (
        "Tu es un assistant de rédaction académique expert, spécialisé dans les mémoires de "
        "licence, master et rapports de stage pour les étudiants d'Afrique francophone. "
        "Tu réponds toujours en français, de façon claire, rigoureuse et structurée. Tu aides à "
        "la problématique, au plan, à la méthodologie, à la revue de littérature et aux citations "
        "(APA, Harvard). Tu ne fabriques jamais de fausses sources."
    ),
    "improve": (
        "Tu es un correcteur académique. Améliore le texte fourni : style, clarté, cohérence, "
        "orthographe et registre académique, en gardant le sens original. Réponds en français et "
        "ne renvoie que le texte amélioré, sans commentaire."
    ),
    "outline": (
        "Tu es un directeur de recherche. À partir du sujet fourni, propose un plan de mémoire "
        "détaillé et hiérarchisé (parties, chapitres, sections) en français, prêt à l'emploi."
    ),
    "expand": (
        "Tu es un rédacteur académique. Développe et enrichis le passage fourni en un texte "
        "académique fluide et argumenté, en français, en conservant le fil directeur."
    ),
}


@api.post("/ai/chat")
async def ai_chat(data: AIChatInput, user: dict = Depends(get_current_user)):
    from emergentintegrations.llm.chat import LlmChat, UserMessage, TextDelta, StreamDone

    mode = data.mode if data.mode in SYSTEM_PROMPTS else "chat"
    session_id = data.document_id or f"chat-{user['id']}"
    prompt_text = data.message
    if data.context:
        prompt_text = f"Contexte du document :\n\"\"\"\n{data.context[:6000]}\n\"\"\"\n\nDemande : {data.message}"

    chat = LlmChat(
        api_key=EMERGENT_LLM_KEY, session_id=session_id, system_message=SYSTEM_PROMPTS[mode],
    ).with_model("anthropic", "claude-sonnet-4-6")

    full = {"text": ""}

    async def gen():
        try:
            async for event in chat.stream_message(UserMessage(text=prompt_text)):
                if isinstance(event, TextDelta):
                    full["text"] += event.content
                    yield _sse(event.content)
                elif isinstance(event, StreamDone):
                    break
        except Exception as e:
            yield _sse(f"[Erreur IA : {e}]")
        finally:
            await db.ai_messages.insert_many([
                {"id": str(uuid.uuid4()), "user_id": user["id"], "document_id": data.document_id,
                 "role": "user", "content": data.message, "created_at": now_iso()},
                {"id": str(uuid.uuid4()), "user_id": user["id"], "document_id": data.document_id,
                 "role": "assistant", "content": full["text"], "created_at": now_iso()},
            ])
            yield "event: done\ndata: [DONE]\n\n"

    return StreamingResponse(
        gen(), media_type="text/event-stream",
        headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no", "Connection": "keep-alive"},
    )


def _sse(text: str) -> str:
    # Encode arbitrary text as a single SSE data event (base64 avoids newline issues).
    import base64
    return f"data: {base64.b64encode(text.encode('utf-8')).decode('ascii')}\n\n"


@api.get("/ai/history")
async def ai_history(document_id: Optional[str] = Query(None), user: dict = Depends(get_current_user)):
    q = {"user_id": user["id"]}
    if document_id:
        q["document_id"] = document_id
    cursor = db.ai_messages.find(q).sort("created_at", 1).limit(200)
    return [clean(m) async for m in cursor]


@api.get("/")
async def root():
    return {"service": "MémoirePro API", "status": "ok"}


app.include_router(api)


# ---------------------------------------------------------------------------
# Startup
# ---------------------------------------------------------------------------
@app.on_event("startup")
async def startup():
    await db.users.create_index("email", unique=True)
    await db.orders.create_index("id", unique=True)
    await db.documents.create_index("id", unique=True)
    admin_email = os.environ.get("ADMIN_EMAIL", "admin@memoirepro.africa").lower()
    admin_password = os.environ.get("ADMIN_PASSWORD", "admin123")
    existing = await db.users.find_one({"email": admin_email})
    if existing is None:
        await db.users.insert_one({
            "id": str(uuid.uuid4()), "email": admin_email, "password_hash": hash_password(admin_password),
            "name": "Administrateur", "role": "admin", "created_at": now_iso(),
        })
    elif not verify_password(admin_password, existing["password_hash"]):
        await db.users.update_one({"email": admin_email}, {"$set": {"password_hash": hash_password(admin_password)}})
    try:
        init_storage()
    except Exception as e:
        print(f"[storage] init failed: {e}")
